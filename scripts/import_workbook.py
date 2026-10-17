"""One-time import of the league workbook into the site's CSV files.

Usage:
    python scripts/import_workbook.py StarWars_DynastyFF_20261017.xlsm

Writes data/franchises.csv, data/seasons.csv, data/contracts.csv,
data/cap_trades.csv, and data/espn/rosters.json (a bootstrap roster
that the nightly ESPN sync replaces).
"""

import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

# Sheet name -> franchise id. Franchise ids are permanent keys.
FRANCHISES = {
    "Renbarger": "renbarger",
    "Andersen": "andersen",
    "B_McMillan": "b_mcmillan",
    "K_McMillan": "k_mcmillan",
    "Tusa_Chen": "tusa_chen",
    "Turowsky": "turowsky",
    "Armijo_Pendleton": "armijo_pendleton",
    "Talaric_Patch": "talaric_patch",
    "Turner": "turner",
    "Bubernak": "bubernak",
    "Stone": "stone",
    "Sands": "sands",
}

# Short owner labels used in cap-trade text ("Renbarger trades B. McMillan $35 in 2027").
TRADE_LABELS = {
    "renbarger": ["renbarger"],
    "andersen": ["andersen"],
    "b_mcmillan": ["b. mcmillan", "b.mcmillan", "brennan"],
    "k_mcmillan": ["k. mcmillan", "k.mcmillan", "kieran"],
    "tusa_chen": ["tusa", "chen"],
    "turowsky": ["turowsky"],
    "armijo_pendleton": ["armijo", "pendleton"],
    "talaric_patch": ["talaric", "patch"],
    "turner": ["turner"],
    "bubernak": ["bubernak"],
    "stone": ["stone"],
    "sands": ["sands"],
}

# Typo fixes applied on import.
NAME_FIXES = {
    "Malik nabers": "Malik Nabers",
    "Rashod bateman": "Rashod Bateman",
    "Emeka Egubka": "Emeka Egbuka",
    "Emmet Johnson": "Emmett Johnson",
    "Will Armjio": "Will Armijo",
}
POS_FIXES = {"Wr": "WR"}
POS_FILL = {"KC Concepcion": "WR"}


def clean(s):
    if s is None:
        return ""
    s = str(s).replace("\t", " ").strip()
    s = re.sub(r"\s+", " ", s)
    return NAME_FIXES.get(s, s)


def find_franchise(text):
    t = text.lower()
    hits = [fid for fid, labels in TRADE_LABELS.items() if any(l in t for l in labels)]
    return hits


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)  # fonts are still read
    info = wb["LeagueInfo"]
    current_season = int(info["M4"].value)

    franchises, seasons, contracts, trades = [], [], [], []
    rosters = {}

    # Current team names / owners from LeagueInfo L6:M17.
    current, council = {}, set()
    for r in range(6, 18):
        team, owners = clean(info.cell(r, 12).value), clean(info.cell(r, 13).value)
        current[team] = owners
        if info.cell(r, 13).font.b:  # Jedi High Council members are bold
            council.add(team)

    for sheet, fid in FRANCHISES.items():
        ws = wb[sheet]
        team_name = clean(ws["B2"].value)
        franchises.append({
            "id": fid,
            "team_name": team_name,
            "owners": current.get(team_name, ""),
            "council": "yes" if team_name in council else "",
            "espn_team_id": "",
        })

        # Season history W:AC, rows 4-33.
        for r in range(4, 34):
            year, name, owners, place, w, l, t = (ws.cell(r, c).value for c in range(23, 30))
            if not name:
                continue
            place = clean(place)
            m = re.match(r"(\d+)\w*(\**)", place)
            seasons.append({
                "season": int(year),
                "franchise": fid,
                "team_name": clean(name),
                "owners": clean(owners),
                "place": int(m.group(1)),
                "mark": m.group(2),
                "wins": w or 0,
                "losses": l or 0,
                "ties": t or 0,
            })

        # Roster B:M, rows 4-24.
        roster = []
        for r in range(4, 25):
            player = clean(ws.cell(r, 3).value)
            if not player:
                continue
            pos = clean(ws.cell(r, 5).value)
            pos = POS_FIXES.get(pos, pos) or POS_FILL.get(player, "")
            drafted = ws.cell(r, 7).value
            years = ws.cell(r, 10).value
            price = ws.cell(r, 11).value
            roster.append({"name": player, "pos": pos, "espn_id": None})
            if drafted:
                contracts.append({
                    "player": player,
                    "pos": pos,
                    "franchise": fid,
                    "draft_year": int(drafted),
                    "price": int(price or 0),
                    "years": int(years) if years else "",
                    "status": "active",
                    "dead_from": "",
                    "espn_id": "",
                    "note": "",
                })
        rosters[fid] = roster

        # Traded cap: P16:U24. Record each trade once, from the sending side.
        for r in range(16, 25):
            text = clean(ws.cell(r, 16).value)
            amt = ws.cell(r, 20).value
            year = ws.cell(r, 21).value
            if not text or not isinstance(amt, (int, float)) or amt >= 0:
                continue
            others = [f for f in find_franchise(text.split("trades", 1)[1]) if f != fid]
            if len(others) != 1:
                raise SystemExit(f"Cannot parse cap trade on {sheet}: {text!r}")
            trades.append({
                "draft_year": int(year),
                "from": fid,
                "to": others[0],
                "amount": int(-amt),
                "note": text,
            })

    DATA.mkdir(exist_ok=True)
    write_csv(DATA / "franchises.csv", franchises)
    seasons.sort(key=lambda s: (s["season"], s["place"]))
    write_csv(DATA / "seasons.csv", seasons)
    contracts.sort(key=lambda c: (c["franchise"], c["draft_year"], -c["price"]))
    write_csv(DATA / "contracts.csv", contracts)
    write_csv(DATA / "cap_trades.csv", trades)

    (DATA / "espn").mkdir(exist_ok=True)
    (DATA / "espn" / "rosters.json").write_text(json.dumps({
        "source": "workbook",
        "season": current_season,
        "updated": re.sub(r"(\d{4})(\d\d)(\d\d)", r"\1-\2-\3", Path(path).stem.split("_")[-1]),
        "teams": rosters,
    }, indent=1))

    print(f"{len(franchises)} franchises, {len(seasons)} seasons, "
          f"{len(contracts)} contracts, {len(trades)} cap trades")


def write_csv(path, rows):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


if __name__ == "__main__":
    main(sys.argv[1])
